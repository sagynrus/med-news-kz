"""Публикует ночной разбор для приложения: razbor.xlsx -> docs/razbor.json.

python3 collector/razbor_json.py /mnt/project-files/med-news/razbor/razbor.xlsx docs/razbor.json
"""
import json, sys
from datetime import datetime, timezone
from openpyxl import load_workbook

src, dst = sys.argv[1], sys.argv[2]
items = []
for r in load_workbook(src, read_only=True).active.iter_rows(min_row=2, values_only=True):
    if not r or not r[3]:
        continue
    razbor, n, lab, title, source, date, link, meaning = (list(r) + [None] * 8)[:8]
    items.append({
        "razbor": razbor.date().isoformat() if isinstance(razbor, datetime) else str(razbor or ""),
        "n": n, "lab": bool(lab), "title": title, "source": source or "",
        "date": date.isoformat(timespec="minutes") if isinstance(date, datetime) else (date or ""),
        "link": link or "", "meaning": meaning or "",
    })
out = {"updated": datetime.now(timezone.utc).isoformat(timespec="seconds"), "items": items}
with open(dst, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print(f"razbor.json: {len(items)} новостей")
